"""Author tam/internal/reportout/sheet.xltx.

Run from this directory, with openpyxl available:

    python make_sheet_xltx.py sheet.xltx

The committed sheet.xltx is the source of the sheet's look, and this is
how it was made. Editing it in Excel is a perfectly good way to restyle
the export; this script is here so the file can also be rebuilt from
scratch, and so its choices can be read as text and diffed.

The template carries the whole look of the exported sheet. Rows 1 to 6 of
column A are a style key: one cell per style the renderer uses, in the order
the Go constants name them. Go reads each anchor's style index, removes the
six rows and then writes the report with those styles. Restyling the export
is an edit to this file, not to Go.
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

SHEET = "Sprint report"

INK = "12263A"      # deep navy, the title
ACCENT = "1C7293"   # teal, section headings and the table header fill
RULE = "C9D6DF"     # hairlines around table cells
QUIET = "5A6B7B"    # caveats

FAMILY = "Calibri"

thin = Side(style="thin", color=RULE)
box = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()
wb.template = True
ws = wb.active
ws.title = SHEET

anchors = [
    ("Title", Font(name=FAMILY, size=16, bold=True, color=INK), None, None, None),
    ("Heading", Font(name=FAMILY, size=13, bold=True, color=ACCENT), None,
     Border(bottom=Side(style="medium", color=ACCENT)), None),
    ("Line", Font(name=FAMILY, size=11, color="1A1A1A"), None, None, None),
    ("Table header", Font(name=FAMILY, size=11, bold=True, color="FFFFFF"),
     PatternFill("solid", fgColor=ACCENT), box,
     Alignment(horizontal="left", vertical="center")),
    ("Table cell", Font(name=FAMILY, size=11, color="1A1A1A"), None, box,
     Alignment(vertical="center")),
    ("Caveat", Font(name=FAMILY, size=10, italic=True, color=QUIET), None, None, None),
]

for i, (label, font, fill, border, align) in enumerate(anchors, start=1):
    c = ws.cell(row=i, column=1, value=label)
    c.font = font
    if fill is not None:
        c.fill = fill
    if border is not None:
        c.border = border
    if align is not None:
        c.alignment = align

ws.column_dimensions["A"].width = 46
for col in "BCDEFGH":
    ws.column_dimensions[col].width = 18
ws.sheet_view.showGridLines = False

wb.save(sys.argv[1])
print("wrote", sys.argv[1])
