"""Author tam/internal/reportout/deck.potx.

Run from this directory, with python-pptx and lxml available:

    python make_deck_potx.py deck.potx

The committed deck.potx is the source of the deck's look, and this is how
it was made. Editing it in PowerPoint is a perfectly good way to restyle
the export; this script is here so the file can also be rebuilt from
scratch, and so its choices can be read as text and diffed.

Two layouts, named so the renderer can find them without counting parts:
"Report title" for the opening slide and "Report section" for every section
slide. The section layout's placeholders are where the renderer puts the
heading, the sentences and the caveats, so their position, size, colour and
font are the template's business and not Go's. The table's look is the
table style in ppt/tableStyles.xml.
"""
import re
import sys
import zipfile

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.util import Emu

from lxml import etree

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"

INK = "12263A"       # deep navy: the title slide, section headings
ACCENT = "1C7293"    # teal: the table's header row
BAND = "EDF3F6"      # the banded row behind every second table row
RULE = "C9D6DF"      # the hairline between table rows
BODY = "3A4A5A"      # the sentences under a heading
QUIET = "5A6B7B"     # the caveats

SLIDE_W = 12192000
SLIDE_H = 6858000
MARGIN = 685800
BODY_W = SLIDE_W - 2 * MARGIN

TITLE_LAYOUT = "Report title"
SECTION_LAYOUT = "Report section"

TABLE_STYLE_ID = "{6E25E649-3F16-4E02-A733-19D2CEDC3F1C}"


def drop(ph):
    ph._element.getparent().remove(ph._element)


def keep_only(layout, idxs):
    for ph in list(layout.placeholders):
        if ph.placeholder_format.idx not in idxs:
            drop(ph)


def place(ph, x, y, cx, cy):
    ph.left, ph.top, ph.width, ph.height = Emu(x), Emu(y), Emu(cx), Emu(cy)


def set_list_style(ph, xml):
    """Replace the placeholder's a:lstStyle, which is what a slide using it
    inherits from."""
    body = ph._element.find(qn("p:txBody"))
    old = body.find(qn("a:lstStyle"))
    if old is not None:
        body.remove(old)
    new = etree.fromstring(xml.encode())
    body.insert(list(body).index(body.find(qn("a:bodyPr"))) + 1, new)


def level(size, colour, bold=False, italic=False, bullet=True, space_after=0, align="l"):
    b = ' b="1"' if bold else ""
    i = ' i="1"' if italic else ""
    no_bullet = "" if bullet else '<a:buNone/>'
    marl = "" if bullet else ' marL="0" indent="0"'
    return (
        f'<a:lstStyle xmlns:a="{A}">'
        f'<a:lvl1pPr{marl} algn="{align}"><a:spcAft><a:spcPts val="{space_after}"/></a:spcAft>{no_bullet}'
        f'<a:defRPr sz="{size}"{b}{i}>'
        f'<a:solidFill><a:srgbClr val="{colour}"/></a:solidFill>'
        f'<a:latin typeface="+mn-lt"/></a:defRPr></a:lvl1pPr></a:lstStyle>'
    )


def solid(element_tag, colour):
    return f'<a:{element_tag}><a:solidFill><a:srgbClr val="{colour}"/></a:solidFill></a:{element_tag}>'


def hairline(tag, colour, width):
    return (f'<a:{tag}><a:ln w="{width}" cap="flat" cmpd="sng" algn="ctr">'
            f'<a:solidFill><a:srgbClr val="{colour}"/></a:solidFill>'
            f'<a:prstDash val="solid"/></a:ln></a:{tag}>')


def no_line(tag):
    return f'<a:{tag}><a:ln><a:noFill/></a:ln></a:{tag}>'


FONT = '<a:font><a:latin typeface="+mn-lt"/><a:ea typeface="+mn-ea"/><a:cs typeface="+mn-cs"/></a:font>'

TABLE_STYLES = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    f'<a:tblStyleLst xmlns:a="{A}" def="{TABLE_STYLE_ID}">'
    f'<a:tblStyle styleId="{TABLE_STYLE_ID}" styleName="Sprint report">'
    '<a:wholeTbl>'
    f'<a:tcTxStyle>{FONT}<a:srgbClr val="1A1A1A"/></a:tcTxStyle>'
    '<a:tcStyle><a:tcBdr>'
    f'{no_line("left")}{no_line("right")}'
    f'{hairline("top", RULE, 12700)}{hairline("bottom", RULE, 12700)}'
    f'{hairline("insideH", RULE, 12700)}{no_line("insideV")}'
    f'</a:tcBdr>{solid("fill", "FFFFFF")}</a:tcStyle>'
    '</a:wholeTbl>'
    f'<a:band1H><a:tcStyle><a:tcBdr/>{solid("fill", BAND)}</a:tcStyle></a:band1H>'
    '<a:firstRow>'
    f'<a:tcTxStyle b="on">{FONT}<a:srgbClr val="FFFFFF"/></a:tcTxStyle>'
    f'<a:tcStyle><a:tcBdr>{hairline("bottom", INK, 19050)}</a:tcBdr>{solid("fill", ACCENT)}</a:tcStyle>'
    '</a:firstRow>'
    '</a:tblStyle>'
    '</a:tblStyleLst>'
)


def background(slide_part, colour):
    cSld = slide_part._element.find(qn("p:cSld"))
    old = cSld.find(qn("p:bg"))
    if old is not None:
        cSld.remove(old)
    bg = etree.fromstring(
        f'<p:bg xmlns:p="{P}" xmlns:a="{A}"><p:bgPr>'
        f'<a:solidFill><a:srgbClr val="{colour}"/></a:solidFill>'
        f'<a:effectLst/></p:bgPr></p:bg>'.encode())
    cSld.insert(0, bg)


def main(out):
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H

    master = prs.slide_master
    layouts = list(master.slide_layouts)
    title_layout, section_layout = layouts[0], layouts[3]
    for layout in layouts:
        if layout not in (title_layout, section_layout):
            master.slide_layouts.remove(layout)

    title_layout.name = TITLE_LAYOUT
    keep_only(title_layout, {0})
    background(title_layout, INK)
    ph = title_layout.placeholders[0]
    place(ph, MARGIN, 2514600, BODY_W, 1600200)
    ph.text_frame.word_wrap = True
    ph.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    ph.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    set_list_style(ph, level(4000, "FFFFFF", bold=True, bullet=False, align="ctr"))

    section_layout.name = SECTION_LAYOUT
    keep_only(section_layout, {0, 1, 2})
    background(section_layout, "FFFFFF")

    heading = section_layout.placeholders[0]
    place(heading, MARGIN, 457200, BODY_W, 838200)
    heading.text_frame.word_wrap = True
    heading.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    heading.text_frame.vertical_anchor = MSO_ANCHOR.BOTTOM
    set_list_style(heading, level(2800, INK, bold=True, bullet=False))

    lines = section_layout.placeholders[1]
    place(lines, MARGIN, 1447800, BODY_W, 685800)
    lines.text_frame.word_wrap = True
    lines.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    lines.text_frame.vertical_anchor = MSO_ANCHOR.TOP
    set_list_style(lines, level(1400, BODY, bullet=False, space_after=300))

    notes = section_layout.placeholders[2]
    place(notes, MARGIN, 5638800, BODY_W, 914400)
    notes.text_frame.word_wrap = True
    notes.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    notes.text_frame.vertical_anchor = MSO_ANCHOR.TOP
    set_list_style(notes, level(1000, QUIET, italic=True, bullet=False, space_after=200))

    background(master, "FFFFFF")

    tmp = out + ".tmp.pptx"
    prs.save(tmp)
    repack(tmp, out)
    print("wrote", out)


def repack(src, out):
    """python-pptx writes a presentation; a template differs in its main
    part's content type. The theme's two colours and the table style are
    written here because python-pptx models neither."""
    zin = zipfile.ZipFile(src)
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            if item.filename == "docProps/thumbnail.jpeg":
                continue
            if item.filename.startswith("ppt/printerSettings/"):
                continue
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                text = data.decode()
                text = text.replace("presentationml.presentation.main+xml",
                                    "presentationml.template.main+xml")
                text = re.sub(r'<Override PartName="/docProps/thumbnail\.jpeg"[^>]*/>', "", text)
                text = re.sub(r'<Default Extension="jpeg"[^>]*/>', "", text)
                text = re.sub(r'<Default Extension="bin"[^>]*/>', "", text)
                data = text.encode()
            elif item.filename == "ppt/presentation.xml":
                data = data.decode().replace('type="screen4x3"', 'type="screen16x9"').encode()
            elif item.filename == "ppt/theme/theme1.xml":
                text = data.decode("utf-8")
                text = text.replace('<a:dk2><a:srgbClr val="1F497D"/></a:dk2>',
                                    f'<a:dk2><a:srgbClr val="{INK}"/></a:dk2>')
                text = text.replace('<a:accent1><a:srgbClr val="4F81BD"/></a:accent1>',
                                    f'<a:accent1><a:srgbClr val="{ACCENT}"/></a:accent1>')
                data = text.encode("utf-8")
            elif item.filename == "ppt/tableStyles.xml":
                data = TABLE_STYLES.encode()
            elif item.filename.endswith(".rels"):
                text = data.decode()
                text = re.sub(r'<Relationship [^>]*printerSettings[^>]*/>', "", text)
                text = re.sub(r'<Relationship [^>]*thumbnail[^>]*/>', "", text)
                data = text.encode()
            zout.writestr(item.filename, data)
    zin.close()


if __name__ == "__main__":
    main(sys.argv[1])
